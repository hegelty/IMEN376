from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_THEME_COLOR

PKG = Path('/home/hegelty/programming/IMEN376/final_presentation_ai_package')
FIG = PKG / 'figures'
OUT = PKG / 'AI_forecasting_6page_ppt_draft.pptx'

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Palette: clean business navy + beer gold + light background
NAVY = RGBColor(15, 35, 64)
BLUE = RGBColor(31, 119, 180)
GOLD = RGBColor(242, 172, 38)
LIGHT = RGBColor(247, 249, 252)
GRAY = RGBColor(92, 103, 116)
DARK = RGBColor(30, 41, 59)
WHITE = RGBColor(255, 255, 255)
GREEN = RGBColor(35, 145, 95)
RED = RGBColor(210, 65, 65)

FONT = 'Aptos'
FONT_KR = 'Malgun Gothic'


def blank_slide():
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = LIGHT
    return slide


def add_top_bar(slide, section='AI Demand Forecasting'):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.42))
    bar.fill.solid(); bar.fill.fore_color.rgb = NAVY
    bar.line.color.rgb = NAVY
    tx = slide.shapes.add_textbox(Inches(0.45), Inches(0.095), Inches(5.5), Inches(0.25))
    p = tx.text_frame.paragraphs[0]
    p.text = section
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = FONT


def add_title(slide, title, subtitle=None, y=0.68):
    tb = slide.shapes.add_textbox(Inches(0.55), Inches(y), Inches(12.1), Inches(0.55))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.font.name = FONT_KR
    if subtitle:
        st = slide.shapes.add_textbox(Inches(0.58), Inches(y+0.62), Inches(11.8), Inches(0.35))
        sp = st.text_frame.paragraphs[0]
        sp.text = subtitle
        sp.font.size = Pt(12)
        sp.font.color.rgb = GRAY
        sp.font.name = FONT_KR


def add_footer(slide, page):
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(7.18), Inches(12.15), Inches(0.012))
    line.fill.solid(); line.fill.fore_color.rgb = RGBColor(215, 221, 230)
    line.line.color.rgb = RGBColor(215, 221, 230)
    tx = slide.shapes.add_textbox(Inches(0.58), Inches(7.22), Inches(7), Inches(0.22))
    p = tx.text_frame.paragraphs[0]
    p.text = 'OB Cass Demand Forecasting | POM Term Project'
    p.font.size = Pt(7.5)
    p.font.color.rgb = GRAY
    p.font.name = FONT
    pg = slide.shapes.add_textbox(Inches(12.0), Inches(7.22), Inches(0.7), Inches(0.22))
    pp = pg.text_frame.paragraphs[0]
    pp.text = f'{page}/6'
    pp.font.size = Pt(7.5)
    pp.font.color.rgb = GRAY
    pp.alignment = PP_ALIGN.RIGHT


def add_card(slide, x, y, w, h, title, body, accent=BLUE, title_size=13, body_size=10):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    card.fill.solid(); card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = RGBColor(226, 232, 240)
    # accent strip
    strip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h))
    strip.fill.solid(); strip.fill.fore_color.rgb = accent
    strip.line.color.rgb = accent
    tx = slide.shapes.add_textbox(Inches(x+0.22), Inches(y+0.15), Inches(w-0.35), Inches(h-0.2))
    tf = tx.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(title_size)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.font.name = FONT_KR
    p.space_after = Pt(5)
    for line in body.split('\n'):
        bp = tf.add_paragraph()
        bp.text = line
        bp.font.size = Pt(body_size)
        bp.font.color.rgb = DARK
        bp.font.name = FONT_KR
        bp.line_spacing = 1.08
    return card


def add_badge(slide, x, y, text, color=BLUE, w=1.1):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.34))
    shape.fill.solid(); shape.fill.fore_color.rgb = color
    shape.line.color.rgb = color
    tx = shape.text_frame
    tx.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tx.paragraphs[0]
    p.text = text
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = FONT_KR
    p.alignment = PP_ALIGN.CENTER
    return shape


def add_image(slide, filename, x, y, w=None, h=None):
    path = FIG / filename
    if w and h:
        return slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
    if w:
        return slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w))
    if h:
        return slide.shapes.add_picture(str(path), Inches(x), Inches(y), height=Inches(h))
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y))


def add_bullets(slide, x, y, w, h, bullets, size=12, color=DARK):
    tx = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tx.text_frame
    tf.clear()
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = b
        p.level = 0
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.font.name = FONT_KR
        p.line_spacing = 1.1
        p.space_after = Pt(5)
    return tx

# Slide 1
slide = blank_slide(); add_top_bar(slide); add_footer(slide, 1)
# hero panel
hero = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(0.42), prs.slide_width, Inches(2.0))
hero.fill.solid(); hero.fill.fore_color.rgb = NAVY
hero.line.color.rgb = NAVY
htext = slide.shapes.add_textbox(Inches(0.65), Inches(0.74), Inches(8.8), Inches(0.7))
p = htext.text_frame.paragraphs[0]
p.text = 'AI 기반 Cass 수요예측 모델'
p.font.size = Pt(32); p.font.bold = True; p.font.color.rgb = WHITE; p.font.name = FONT_KR
st = slide.shapes.add_textbox(Inches(0.68), Inches(1.47), Inches(9.3), Inches(0.52))
sp = st.text_frame.paragraphs[0]
sp.text = '국내 맥주 apparent consumption 예측 → 발효탱크 사전 배분 의사결정'
sp.font.size = Pt(15); sp.font.color.rgb = RGBColor(222,235,255); sp.font.name = FONT_KR
add_badge(slide, 10.35, 0.95, 'Chronos-2', GOLD, 1.35)
add_badge(slide, 11.85, 0.95, 'MVP', GREEN, 0.75)
add_card(slide, 0.72, 2.82, 3.75, 2.0, 'Problem', '발효 리드타임 때문에\n수요 발생 후 즉시 대응 불가\n→ 사전 예측이 생산계획의 핵심', RED)
add_card(slide, 4.72, 2.82, 3.75, 2.0, 'Target', '국내 출고량 + HS2203 수입 − 수출\n= 국내 소비 가능 물량 proxy\n논알콜/무알콜 세그먼트 포함', BLUE)
add_card(slide, 8.72, 2.82, 3.75, 2.0, 'Result', 'Chronos-2가 최종 1위\nMASE 0.387 / WAPE 0.0168\nSeasonal Naive 대비 약 13% 개선', GREEN)
add_image(slide, '08_final_model_summary_table.png', 1.35, 4.98, w=10.4)

# Slide 2
slide = blank_slide(); add_top_bar(slide); add_title(slide, '1. 수요 Target: 국내 소비 가능 물량으로 재정의', '단순 국내 출고량이 아니라 수입 유입과 수출 차감을 반영한 apparent consumption proxy'); add_footer(slide, 2)
add_image(slide, '01_target_construction_apparent_consumption.png', 0.65, 1.55, w=7.35)
add_card(slide, 8.25, 1.55, 4.4, 1.3, 'Target Formula', 'Demandₜ = Domestic Shipmentₜ\n+ Import(HS2203)ₜ − Export(HS2203)ₜ', BLUE, title_size=13, body_size=11)
add_card(slide, 8.25, 3.05, 4.4, 1.4, 'Why this matters', '국내 생산/출고 중심 지표를\n국내에서 실제 소비 가능한 물량에\n더 가까운 target으로 개선', GOLD, title_size=13, body_size=10.5)
add_card(slide, 8.25, 4.65, 4.4, 1.48, 'Scope note', '실제 OB 내부 SKU 수요는 아님\n공개자료 기반 diagnostic proxy\n→ Phase 0 MVP로 해석', RED, title_size=13, body_size=10.5)
add_bullets(slide, 0.82, 6.35, 11.5, 0.45, ['Non-alcohol / zero-alcohol segment는 Cass 0.0 성장 SKU 시나리오 보조 layer로 별도 반영'], size=11, color=NAVY)

# Slide 3
slide = blank_slide(); add_top_bar(slide); add_title(slide, '2. Feature 설계: 수요를 움직이는 외생변수', '날씨·스포츠·달력·가격·뉴스·검색량을 구분해 leakage 없이 사용'); add_footer(slide, 3)
add_card(slide, 0.65, 1.72, 2.4, 1.25, 'Weather', '월평균기온\n폭염일수 / 열대야\n강수량·비 오는 날', BLUE)
add_card(slide, 3.25, 1.72, 2.4, 1.25, 'Sports & Events', 'KBO 경기수\n월드컵/대형 이벤트\n지역·맥주 축제', GOLD)
add_card(slide, 5.85, 1.72, 2.4, 1.25, 'Calendar', '공휴일 수\n주말·연휴\n평일 공휴일', GREEN)
add_card(slide, 8.45, 1.72, 2.4, 1.25, 'Price / Economy', '수입맥주 단가 YoY\nCPI 상승률\nOB 가격인상', RED)
add_card(slide, 0.65, 3.35, 3.15, 1.48, 'Google Trends rule', 'Target 생성에는 미사용\n예측에는 전월(t-1) 검색량만 테스트\n→ demand-sensing 보조 신호', NAVY, body_size=10)
add_card(slide, 4.05, 3.35, 3.15, 1.48, 'News / Sentiment', '맥주·Cass·비알콜 뉴스\n날씨-수요 기사\n축제/행사 뉴스', BLUE, body_size=10)
add_card(slide, 7.45, 3.35, 4.1, 1.48, 'Model candidates', 'Chronos-2, Seasonal Naive, Ridge/Lasso/ElasticNet, RandomForest, ExtraTrees, GradientBoosting, HistGB, KNN, SVR', GOLD, body_size=10)
add_card(slide, 1.0, 5.35, 10.8, 0.95, 'Validation rule', 'Rolling backtest: 2024–2025 각 월 예측 시 해당 월 이전 데이터만 학습  |  평가지표: MASE / WAPE / RMSE — 낮을수록 좋음', BLUE, title_size=12, body_size=10.2)

# Slide 4
slide = blank_slide(); add_top_bar(slide); add_title(slide, '3. 모델 비교 결과: Chronos-2가 최종 1위', '국내 맥주 apparent consumption target 기준으로 다양한 AI/ML 모델 비교'); add_footer(slide, 4)
add_image(slide, '03_model_comparison_mase.png', 0.6, 1.98, w=6.0)
add_image(slide, '04_model_comparison_wape.png', 6.92, 1.98, w=5.8)
add_card(slide, 0.8, 6.28, 3.8, 0.7, 'Key result', 'Chronos-2 MASE 0.387 / WAPE 0.0168', GREEN, title_size=11, body_size=9.5)
add_card(slide, 4.85, 6.28, 3.8, 0.7, 'Baseline', 'Seasonal Naive MASE 0.445 / WAPE 0.0193', BLUE, title_size=11, body_size=9.5)
add_card(slide, 8.9, 6.28, 3.4, 0.7, 'Improvement', '약 13% error reduction', GOLD, title_size=11, body_size=9.5)

# Slide 5
slide = blank_slide(); add_top_bar(slide); add_title(slide, '4. Backtest: 예측 흐름과 월별 오차', 'Chronos-2는 강한 baseline을 개선하면서 예측구간까지 제공'); add_footer(slide, 5)
add_image(slide, '02_backtest_actual_vs_chronos2_seasonal.png', 0.65, 1.78, w=7.55)
add_card(slide, 0.9, 5.78, 7.05, 0.65, 'Error view', '월별 절대오차 비교 그래프는 package figures/05에 포함 — 발표 본문에서는 예측 흐름 중심으로 제시', BLUE, title_size=12, body_size=10)
add_card(slide, 8.45, 1.82, 4.0, 1.25, 'Interpretation', 'Seasonal Naive도 강한 기준선이지만\nChronos-2가 전체 테스트에서\nMASE/WAPE 모두 개선', GREEN, body_size=10.2)
add_card(slide, 8.45, 3.28, 4.0, 1.25, 'Why interval matters', '발효탱크 운영은 평균보다\n상한 리스크가 중요\n→ Q90 기반 안전계획 검토', GOLD, body_size=10.2)
add_card(slide, 8.45, 4.74, 4.0, 1.25, 'Google Trends', '전달 검색량은 일부 ML에서 개선\n하지만 핵심 모델은 Chronos-2\nGT는 보조 demand-sensing 신호', BLUE, body_size=10.2)

# Slide 6
slide = blank_slide(); add_top_bar(slide); add_title(slide, '5. 운영 적용: 예측을 발효탱크 계획으로 연결', 'AI forecast → baseline check → risk overlay → production planning'); add_footer(slide, 6)
# workflow boxes
steps = [
    ('Chronos-2\n월별 예측', BLUE),
    ('Seasonal Naive\nSanity Check', GOLD),
    ('Q90 상한\nRisk Overlay', RED),
    ('발효탱크·안전재고\n계획 검토', GREEN),
]
x0 = 0.75
for i, (txt, col) in enumerate(steps):
    x = x0 + i*3.1
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.65), Inches(2.45), Inches(1.05))
    box.fill.solid(); box.fill.fore_color.rgb = col
    box.line.color.rgb = col
    tf = box.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.text = txt; p.font.size = Pt(14); p.font.bold = True; p.font.color.rgb = WHITE; p.font.name = FONT_KR; p.alignment = PP_ALIGN.CENTER
    if i < len(steps)-1:
        arr = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x+2.48), Inches(1.92), Inches(0.58), Inches(0.45))
        arr.fill.solid(); arr.fill.fore_color.rgb = RGBColor(148,163,184)
        arr.line.color.rgb = RGBColor(148,163,184)
add_image(slide, '07_event_stress_months_on_target.png', 0.7, 3.15, w=5.95)
add_image(slide, '06_nonalc_zero_segment_proxy.png', 6.7, 3.15, w=5.95)


prs.save(OUT)
print(OUT)
